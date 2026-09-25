# -*- coding: utf-8 -*-
"""
Card-aware fit check.

PowerPoint text boxes auto-grow, so a declared height of 0.2" means nothing.
What matters is whether the GROWN text still sits inside the card rectangle
behind it. This measures each text box against the nearest enclosing card.
"""
from pptx import Presentation

EMU = 914400.0
CHAR_W = 0.48   # Calibri average advance as a fraction of point size
LINE_H = 1.22   # line box as a fraction of point size


def needed_height(tf, w_in):
    total = 0.0
    for para in tf.paragraphs:
        txt = "".join(r.text for r in para.runs)
        if not txt:
            continue
        sz = max((r.font.size.pt for r in para.runs if r.font.size), default=12)
        per_line = max(1, int(w_in / (CHAR_W * sz / 72.0)))
        lines = max(1, -(-len(txt) // per_line))
        ls = para.line_spacing if isinstance(para.line_spacing, float) else 1.0
        total += lines * (LINE_H * sz / 72.0) * ls
        total += (para.space_after.pt / 72.0) if para.space_after else 0.0
    return total


def main():
    prs = Presentation("InstantZero-final.pptx")
    worst = []

    for idx in (7, 8, 9, 11):
        s = prs.slides[idx - 1]

        # Cards are autoshape rectangles taller than 1.2" — but EXCLUDE the
        # full-slide background rect, which would otherwise match every text
        # box and make everything look like it has acres of room.
        cards = []
        for sh in s.shapes:
            w_, h_ = sh.width / EMU, sh.height / EMU
            if sh.shape_type != 1 or h_ <= 1.2 or w_ <= 2.0:
                continue
            if w_ > 12.5 and h_ > 7.0:          # background
                continue
            cards.append((sh.left / EMU, sh.top / EMU, w_, h_))

        for sh in s.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            x, y = sh.left / EMU, sh.top / EMU
            w = sh.width / EMU
            need = needed_height(sh.text_frame, w)
            bottom = y + need

            # Smallest enclosing card wins, so nested rectangles cannot
            # inflate the available space.
            host = None
            for (cx, cy, cw, ch) in cards:
                if cx - 0.02 <= x and x + w <= cx + cw + 0.02 and cy <= y <= cy + ch:
                    if host is None or (cw * ch) < (host[2] * host[3]):
                        host = (cx, cy, cw, ch)
            if not host:
                continue

            limit = host[1] + host[3] - 0.10          # 0.10" bottom padding
            slack = limit - bottom
            label = sh.text_frame.text.strip().replace("\n", " / ")[:52]
            worst.append((slack, idx, round(need, 2), round(limit - y, 2), label))

    worst.sort()
    print("%-6s %-5s %-7s %-9s %s" % ("SLACK", "SLIDE", "NEEDS", "AVAIL", "TEXT"))
    for slack, idx, need, avail, label in worst[:12]:
        flag = "OVERFLOW" if slack < 0 else ("tight" if slack < 0.15 else "ok")
        print("%-6.2f %-5d %-7.2f %-9.2f %-8s %s" % (slack, idx, need, avail, flag, label))

    bad = [w for w in worst if w[0] < 0.15]
    print()
    print("blocks with < 0.15\" slack:", len(bad))
    return 1 if any(w[0] < 0 for w in worst) else 0


if __name__ == "__main__":
    raise SystemExit(main())
