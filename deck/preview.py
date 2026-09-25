# -*- coding: utf-8 -*-
"""
Approximate renderer for visual QA.

No LibreOffice here, so this draws each slide's shapes and text with Pillow at
the real aspect ratio. Fonts and wrapping are approximations — this is for
spotting layout faults (overlaps, collisions, bad gaps, stray shapes), not for
judging final typography.
"""
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

EMU = 914400.0
SCALE = 110          # px per inch
W, H = int(13.333 * SCALE), int(7.5 * SCALE)


def load_font(size_px, bold=False):
    names = (["arialbd.ttf", "calibrib.ttf"] if bold else ["arial.ttf", "calibri.ttf"])
    for n in names:
        try:
            return ImageFont.truetype(n, size_px)
        except Exception:
            continue
    return ImageFont.load_default()


def rgb_of(shape):
    try:
        if shape.fill.type is not None and shape.fill.type == 1:
            c = shape.fill.fore_color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return None


def wrap(draw, txt, font, max_px):
    words, lines, cur = txt.split(), [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textlength(trial, font=font) <= max_px or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def render(prs, index, out):
    slide = prs.slides[index - 1]
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)

    for sh in slide.shapes:
        x, y = sh.left / EMU * SCALE, sh.top / EMU * SCALE
        w, h = sh.width / EMU * SCALE, sh.height / EMU * SCALE

        fill = rgb_of(sh)
        if fill is not None and not sh.has_text_frame:
            d.rectangle([x, y, x + w, y + h], fill=fill)
        elif fill is not None:
            d.rectangle([x, y, x + w, y + h], fill=fill)

        if sh.shape_type == 1 and fill == (255, 255, 255):
            d.rectangle([x, y, x + w, y + h], outline=(214, 223, 234), width=1)

        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue

        cy = y
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs)
            if not t:
                continue
            r0 = para.runs[0]
            pt = r0.font.size.pt if r0.font.size else 12
            bold = bool(r0.font.bold)
            try:
                col = r0.font.color.rgb
                col = (col[0], col[1], col[2])
            except Exception:
                col = (20, 20, 20)
            f = load_font(max(8, int(pt / 72.0 * SCALE)), bold)
            for line in wrap(d, t, f, max(10, w)):
                d.text((x, cy), line, font=f, fill=col)
                cy += int(pt / 72.0 * SCALE * 1.22)
            cy += int((para.space_after.pt if para.space_after else 0) / 72.0 * SCALE)

    img.save(out, quality=92)
    print("wrote", out)


if __name__ == "__main__":
    prs = Presentation("InstantZero-final.pptx")
    for i in (7, 8, 9, 12):
        render(prs, i, "preview-%d.jpg" % i)
